"""
Excel 读写 — 读模板、写评分、保存结果。
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment

from outputs.config import (
    DIM_ROW,
    EXCEL_SHEET_NAME,
    TASK_TO_COLUMN,
)
from domain.models import TaskScore
from domain.rules_v5 import RULES_TEXT


def fill_excel(
    excel_path: Path,
    output_path: Path,
    tasks: dict[int, TaskScore],
    *,
    metadata: dict | None = None,
) -> None:
    """将评分、审计批注和运行元数据写入 Excel。"""
    wb = openpyxl.load_workbook(str(excel_path))
    ws = wb[EXCEL_SHEET_NAME]

    # 写入评分
    for task_idx, task_score in sorted(tasks.items()):
        col_idx = TASK_TO_COLUMN.get(task_idx)
        if col_idx is None:
            print(f"  ⚠ 跳过 task {task_idx}：无对应 Excel 列", file=sys.stderr)
            continue

        for dim_name, result in task_score.stages.items():
            row = DIM_ROW.get(dim_name)
            if row is None:
                continue
            cell = ws.cell(row, col_idx)
            cell.value = result.score
            lines = [
                f"规则版本: {result.rule_version}",
                f"置信度: {result.confidence:.0%}",
                f"需人工复核: {'是' if result.needs_review else '否'}",
            ]
            if result.review_reason:
                lines.append(f"复核原因: {result.review_reason}")
            if result.missing_evidence:
                lines.append("缺失证据: " + "、".join(result.missing_evidence))
            lines.extend(result.evidence)
            cell.comment = Comment("\n".join(lines), "inference_scorer")

    if metadata:
        n_action_steps = metadata.get("n_action_steps")
        if n_action_steps is not None:
            ws["C28"] = f"N_ACTION_STEPS_OVERRIDE = {n_action_steps}"
        num_inference_steps = metadata.get("num_inference_steps")
        ws["C29"] = (
            "NUM_INFERENCE_STEPS_OVERRIDE = "
            f"{num_inference_steps if num_inference_steps is not None else 'None'}"
        )
        prompt = metadata.get("task_prompt")
        if prompt:
            ws["C30"] = f"提示词：{prompt}"
        video_path = metadata.get("video_path")
        model_path = metadata.get("ckpt_dir")
        paths = []
        if video_path:
            paths.append(f"video路径：{video_path}")
        if model_path:
            paths.append(f"模型路径：{model_path}")
        if paths:
            ws["C31"] = "\n".join(paths)

    # openpyxl 不计算公式，要求 Excel/WPS 打开时强制重算总分和进度。
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"

    wb.save(str(output_path))
    print(f"  已保存: {output_path}")


def print_rules() -> None:
    """打印评分规则。"""
    print(RULES_TEXT)
